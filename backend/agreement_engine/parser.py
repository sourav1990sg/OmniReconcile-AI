"""Document text extraction for commercial agreements (PDF / DOCX / OCR fallback)."""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    filename: str
    text: str
    method: str  # pdf_text | docx | ocr | plain | empty
    page_count: int = 0
    ocr_used: bool = False
    notes: list[str] | None = None

    def __post_init__(self) -> None:
        if self.notes is None:
            self.notes = []


class AgreementParser:
    """
    Accept PDF, DOCX, scanned PDF, email PDF, amendments, pricing sheets.

    OCR is best-effort when text layer is empty and optional deps are installed.
    Never invents content — empty text yields empty ParsedDocument.
    """

    SUPPORTED_SUFFIXES = (".pdf", ".docx", ".doc", ".txt", ".md")

    def parse(self, filename: str, content: bytes) -> ParsedDocument:
        name = (filename or "agreement").strip()
        lower = name.lower()
        notes: list[str] = []

        if lower.endswith(".docx") or lower.endswith(".doc"):
            text, method, n = self._parse_docx(content, notes)
            return ParsedDocument(name, text, method, page_count=n, notes=notes)

        if lower.endswith(".pdf"):
            text, method, pages, ocr = self._parse_pdf(content, notes)
            return ParsedDocument(
                name, text, method, page_count=pages, ocr_used=ocr, notes=notes
            )

        if lower.endswith((".txt", ".md")):
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("latin-1", errors="ignore")
            return ParsedDocument(name, text, "plain", page_count=1, notes=notes)

        notes.append(f"Unsupported extension for '{name}'; attempted as plain text.")
        try:
            text = content.decode("utf-8", errors="ignore")
        except Exception:  # noqa: BLE001
            text = ""
        return ParsedDocument(name, text, "plain", notes=notes)

    def _parse_docx(self, content: bytes, notes: list[str]) -> tuple[str, str, int]:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            notes.append("python-docx not installed; DOCX text unavailable.")
            return "", "empty", 0
        try:
            doc = Document(io.BytesIO(content))
            parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
            return text, "docx", max(1, len(doc.paragraphs) // 40)
        except Exception as exc:  # noqa: BLE001
            notes.append(f"DOCX parse failed: {exc}")
            return "", "empty", 0

    def _parse_pdf(
        self, content: bytes, notes: list[str]
    ) -> tuple[str, str, int, bool]:
        text, pages = self._pdf_text_layer(content, notes)
        if text.strip():
            return text, "pdf_text", pages, False

        ocr_text, ocr_pages = self._pdf_ocr(content, notes)
        if ocr_text.strip():
            return ocr_text, "ocr", ocr_pages or pages, True

        notes.append("PDF had no extractable text and OCR unavailable or failed.")
        return "", "empty", pages, False

    def _pdf_text_layer(self, content: bytes, notes: list[str]) -> tuple[str, int]:
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError:
            notes.append("pypdf not installed; PDF text layer unavailable.")
            return "", 0
        try:
            reader = PdfReader(io.BytesIO(content))
            chunks: list[str] = []
            for page in reader.pages:
                try:
                    chunks.append(page.extract_text() or "")
                except Exception:  # noqa: BLE001
                    chunks.append("")
            return "\n".join(chunks), len(reader.pages)
        except Exception as exc:  # noqa: BLE001
            notes.append(f"PDF text extract failed: {exc}")
            return "", 0

    def _pdf_ocr(self, content: bytes, notes: list[str]) -> tuple[str, int]:
        try:
            import pytesseract  # type: ignore
            from pdf2image import convert_from_bytes  # type: ignore
        except ImportError:
            notes.append("OCR deps (pytesseract/pdf2image) not installed.")
            return "", 0
        try:
            images = convert_from_bytes(content, dpi=200)
            texts = [pytesseract.image_to_string(img) for img in images]
            return "\n".join(texts), len(images)
        except Exception as exc:  # noqa: BLE001
            notes.append(f"OCR failed: {exc}")
            logger.warning("Agreement OCR failed: %s", exc)
            return "", 0
